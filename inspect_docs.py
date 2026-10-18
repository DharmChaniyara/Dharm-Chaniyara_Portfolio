import os
import sys

def read_docx(path):
    try:
        import docx
        doc = docx.Document(path)
        print("=== PARAGRAPHS ===")
        for p in doc.paragraphs:
            if p.text.strip():
                print(p.text)
        print("=== TABLES ===")
        for table in doc.tables:
            for row in table.rows:
                print(" | ".join([cell.text.strip().replace('\n', ' ') for cell in row.cells]))
    except Exception as e:
        print(f"Error reading docx with python-docx: {e}")
        # fallback unzip document.xml
        import zipfile
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(path) as z:
            xml_content = z.read('word/document.xml')
            tree = ET.fromstring(xml_content)
            namespaces = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            text_elems = tree.findall('.//w:t', namespaces)
            texts = [elem.text for elem in text_elems if elem.text]
            print("Fallback docx text:")
            print(" ".join(texts))

def read_pdf(path):
    print(f"\n=== PDF: {path} ===")
    try:
        import pypdf
        reader = pypdf.PdfReader(path)
        for i, page in enumerate(reader.pages):
            print(f"--- Page {i+1} ---")
            print(page.extract_text())
    except Exception as e:
        try:
            import fitz
            doc = fitz.open(path)
            for i, page in enumerate(doc):
                print(f"--- Page {i+1} ---")
                print(page.get_text())
        except Exception as e2:
            print(f"Could not read PDF with pypdf or pymupdf: {e2}")

if __name__ == '__main__':
    base = os.path.dirname(os.path.abspath(__file__))
    docx_file = os.path.join(base, 'Dharm_Chaniyara_CV_Draft.docx')
    if os.path.exists(docx_file):
        read_docx(docx_file)
    for f in os.listdir(base):
        if f.endswith('.pdf'):
            read_pdf(os.path.join(base, f))
